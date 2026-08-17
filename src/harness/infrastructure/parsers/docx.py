"""DOCX parser using python-docx (TASK-009).

Headings become heading parts (with level); other paragraphs become
paragraph parts. Embedded images are extracted from the package's related
parts, fingerprinted and emitted as image parts.
"""

import io

from docx import Document

from harness.domain.artifacts.fetching import FetchedArtifact
from harness.domain.artifacts.parsing import ParsedDocument, ParsedPart, PartType
from harness.infrastructure.parsers.images import image_part


def _heading_level(style_name: str) -> int | None:
    if not style_name.startswith("Heading"):
        return None
    suffix = style_name.removeprefix("Heading").strip()
    return int(suffix) if suffix.isdigit() else 1


class DocxParser:
    """Satisfies DocumentParser."""

    def supports(self, mime_type: str, filename: str) -> bool:
        return (
            mime_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or filename.lower().endswith(".docx")
        )

    async def parse(self, artifact: FetchedArtifact) -> ParsedDocument:
        document = Document(io.BytesIO(artifact.content))
        parts: list[ParsedPart] = []
        ordinal = 0
        title: str | None = None

        for paragraph_index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style is not None else ""
            level = _heading_level(style_name or "")
            if level is not None:
                if title is None:
                    title = text
                parts.append(
                    ParsedPart(
                        part_type=PartType.HEADING,
                        ordinal=ordinal,
                        text_content=text,
                        location={"paragraph": paragraph_index},
                        metadata={"level": level},
                    )
                )
            else:
                parts.append(
                    ParsedPart(
                        part_type=PartType.PARAGRAPH,
                        ordinal=ordinal,
                        text_content=text,
                        location={"paragraph": paragraph_index},
                    )
                )
            ordinal += 1

        for related in document.part.related_parts.values():
            content_type = getattr(related, "content_type", "")
            if content_type.startswith("image/"):
                parts.append(
                    image_part(
                        related.blob,
                        ordinal=ordinal,
                        extra_metadata={"content_type": content_type},
                    )
                )
                ordinal += 1

        return ParsedDocument(parts=parts, title=title)
