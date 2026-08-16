"""PPTX parser using python-pptx with slide boundaries (TASK-010).

Each slide becomes one slide part carrying its concatenated text (titles
first), located by slide number. Pictures are extracted per slide as image
parts sharing the slide's location so provenance survives.
"""

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from harness.domain.artifacts.fetching import FetchedArtifact
from harness.domain.artifacts.parsing import ParsedDocument, ParsedPart, PartType
from harness.infrastructure.parsers.images import image_part

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class PptxParser:
    """Satisfies DocumentParser."""

    def supports(self, mime_type: str, filename: str) -> bool:
        return mime_type == PPTX_MIME or filename.lower().endswith(".pptx")

    async def parse(self, artifact: FetchedArtifact) -> ParsedDocument:
        presentation = Presentation(io.BytesIO(artifact.content))
        parts: list[ParsedPart] = []
        ordinal = 0
        title: str | None = None

        for slide_index, slide in enumerate(presentation.slides):
            slide_number = slide_index + 1
            texts: list[str] = []
            slide_title = slide.shapes.title.text.strip() if slide.shapes.title else ""
            if slide_title:
                texts.append(slide_title)
                if title is None:
                    title = slide_title
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        texts.append(text)
            if texts:
                parts.append(
                    ParsedPart(
                        part_type=PartType.SLIDE,
                        ordinal=ordinal,
                        text_content="\n".join(texts),
                        location={"slide": slide_number},
                        metadata={"title": slide_title or None},
                    )
                )
                ordinal += 1
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    parts.append(
                        image_part(
                            shape.image.blob,
                            ordinal=ordinal,
                            location={"slide": slide_number},
                            extra_metadata={"content_type": shape.image.content_type},
                        )
                    )
                    ordinal += 1

        return ParsedDocument(parts=parts, title=title)
