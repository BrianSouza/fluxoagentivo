"""Synthetic in-memory fixtures for parser tests (16_DOCKER_LOCAL.md §3)."""

import io

import pymupdf
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def make_png(
    color: tuple[int, int, int] = (200, 30, 30), size: tuple[int, int] = (64, 48)
) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format="PNG")
    return buffer.getvalue()


def make_pdf() -> bytes:
    document = pymupdf.open()
    page_one = document.new_page()
    page_one.insert_text((72, 100), "Checkout flow overview")
    page_one.insert_image(pymupdf.Rect(72, 150, 172, 230), stream=make_png())
    page_two = document.new_page()
    page_two.insert_text((72, 100), "Backend services involved")
    data = document.tobytes()
    document.close()
    return bytes(data)


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Checkout architecture", level=1)
    document.add_paragraph("The mobile checkout runs inside a WebView.")
    document.add_heading("Backend", level=2)
    document.add_paragraph("The WebView calls the Checkout API.")
    document.add_picture(io.BytesIO(make_png()), width=Inches(1))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title and content

    slide_one = presentation.slides.add_slide(layout)
    slide_one.shapes.title.text = "Checkout flow"
    slide_one.placeholders[1].text = "Mobile app opens a WebView"

    slide_two = presentation.slides.add_slide(layout)
    slide_two.shapes.title.text = "Services"
    slide_two.placeholders[1].text = "Checkout API and Payment Gateway"
    slide_two.shapes.add_picture(io.BytesIO(make_png()), Inches(1), Inches(3), Inches(2))

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


SAMPLE_HTML = b"""
<html>
  <head><title>Checkout flow</title><style>body { color: red; }</style></head>
  <body>
    <script>console.log("boilerplate");</script>
    <h1>Checkout flow</h1>
    <p>The mobile app opens a <a href="/wiki/webview">WebView</a>.</p>
    <img src="/images/checkout.png" alt="Checkout diagram">
    <h2>Backend services</h2>
    <p>The WebView calls the <a href="/wiki/checkout-api">Checkout API</a>.</p>
    <ul><li>Checkout API</li><li>Payment Gateway</li></ul>
  </body>
</html>
"""
